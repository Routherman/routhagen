import openpyxl, json, re, os

wb = openpyxl.load_workbook('candidatos_eze.xlsx', data_only=True)

def norm_phone(raw):
    if not raw: return ''
    s = re.sub(r'[^\d]', '', str(raw))
    if not s: return ''
    if s.startswith('519') and len(s) >= 12:
        s = '549' + s[3:]
    elif s.startswith('51') and len(s) >= 11 and (s[2:4] in ['11', '35', '34', '26', '29', '38']):
        s = '549' + s[2:]
    elif s.startswith('549') and len(s) >= 12:
        pass
    elif s.startswith('54') and not s.startswith('549') and len(s) >= 11:
        s = '549' + s[2:]
    elif len(s) == 10:
        s = '549' + s
    elif len(s) == 8:
        s = '54911' + s
    return s

def format_dni(val):
    if not val: return ''
    digits = re.sub(r'[^\d]', '', str(val))
    if len(digits) == 8:
        return f"{digits[:2]}.{digits[2:5]}.{digits[5:]}"
    elif len(digits) == 7:
        return f"{digits[:1]}.{digits[1:4]}.{digits[4:]}"
    return str(val).strip()

records = []
seen_phones = set()

# 1. Hojas con datos completos
named_sheets = ['Seleccionados', 'Finalistas', 'Recibidos', 'Descartados']
for sheet in named_sheets:
    if sheet not in wb.sheetnames: continue
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    for r in rows[1:]:
        nombre = str(r[0]).strip() if r[0] and str(r[0]).strip().lower() not in ['sin nombre', '-', 'none'] else ''
        email = str(r[1]).strip().lower() if r[1] and '@' in str(r[1]) else ''
        dni = format_dni(r[2])
        tel = norm_phone(r[3])
        edad = str(r[4] or '').strip()
        zona = str(r[5] or '').strip()
        rol = str(r[6] or '').strip()
        url = str(r[7] or '').strip()

        if not tel and not email:
            continue
        
        if not tel and email:
            tel = 'email_' + re.sub(r'[^a-zA-Z0-9]', '', email)

        records.append({
            'nombre': nombre,
            'email': email,
            'dni': dni,
            'telefono': tel,
            'edad': edad,
            'zona': zona,
            'rol': rol,
            'url_cv': url,
            'hoja_origen': sheet
        })
        if tel: seen_phones.add(tel)

# 2. Hoja ezequiel2 (números de teléfono directos)
if 'ezequiel2' in wb.sheetnames:
    ws_ez = wb['ezequiel2']
    for r in ws_ez.iter_rows(values_only=True):
        raw = r[1] if len(r) > 1 and r[1] else (r[0] if r[0] else None)
        if not raw: continue
        tel = norm_phone(raw)
        if not tel or len(tel) < 9: continue
        if tel in seen_phones: continue

        records.append({
            'nombre': f"Prospecto ({tel})",
            'email': '',
            'dni': '',
            'telefono': tel,
            'edad': '',
            'zona': '',
            'rol': 'Contacto telefónico Ezequiel',
            'url_cv': '',
            'hoja_origen': 'ezequiel2'
        })
        seen_phones.add(tel)

out_path = os.path.join(os.path.dirname(__file__), 'candidatos_extracted.json')
with open(out_path, 'w', encoding='utf-8') as f:
    json.dump(records, f, ensure_ascii=False, indent=2)

print(f"Extracted {len(records)} records successfully to {out_path}")
